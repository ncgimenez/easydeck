import io
import json
import streamlit as st
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pypdf import PdfReader

# Configuration centrée
st.set_page_config(page_title="DuoDeck", layout="centered", page_icon="🪄")

# --- INJECTION CSS POUR DESIGN MODERNE & ÉPURÉ ---
st.markdown("""
<style>
    /* Police et espacements globaux */
    .block-container {
        padding-top: 2.5rem;
        padding-bottom: 4rem;
        max-width: 820px;
    }
    
    /* En-tête principal centré */
    .main-header {
        text-align: center;
        margin-bottom: 2rem;
    }
    .main-title {
        font-size: 2.4rem;
        font-weight: 800;
        letter-spacing: -0.02em;
        color: #1e293b;
        margin-bottom: 0.4rem;
    }
    .main-subtitle {
        font-size: 1.05rem;
        color: #64748b;
        margin-bottom: 1.5rem;
    }
    
    /* Cartes de contenu */
    .custom-card {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 16px;
        padding: 1.5rem;
        margin-bottom: 1.5rem;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -2px rgba(0, 0, 0, 0.05);
    }
    
    /* Carte de prévisualisation de slide */
    .slide-preview-card {
        background-color: #f8fafc;
        border-left: 4px solid #3b82f6;
        border-radius: 10px;
        padding: 1rem 1.25rem;
        margin-bottom: 1rem;
    }
    .slide-title-tag {
        font-weight: 700;
        color: #0f172a;
        font-size: 1.1rem;
        margin-bottom: 0.5rem;
    }
    
    /* Stylisation des boutons Streamlit */
    div.stButton > button:first-child {
        width: 100%;
        border-radius: 12px;
        height: 3.2rem;
        font-size: 1.1rem;
        font-weight: 700;
        letter-spacing: 0.01em;
        box-shadow: 0 4px 14px 0 rgba(37, 99, 235, 0.3);
        transition: all 0.2s ease-in-out;
    }
    div.stButton > button:first-child:hover {
        transform: translateY(-1px);
        box-shadow: 0 6px 20px 0 rgba(37, 99, 235, 0.35);
    }
</style>
""", unsafe_allow_html=True)

# En-tête
st.markdown("""
<div class="main-header">
    <div class="main-title">🪄 DuoDeck</div>
    <div class="main-subtitle">Concevez en un clic des présentations différenciées adaptées à chaque profil d'élève.</div>
</div>
""", unsafe_allow_html=True)

api_key = st.secrets.get("GEMINI_API_KEY", None)

# --- MENU LATÉRAL : ARRIÈRE-PLAN & MODÈLE ---
with st.sidebar:
    st.header("🎨 Thème & Arrière-plan")
    bg_file = st.file_uploader("Fond personnalisé (PNG/JPG)", type=["png", "jpg", "jpeg"])
    
    st.divider()
    st.header("📂 Ancien cours (Optionnel)")
    st.caption("Importez un ancien fichier pour que l'IA imite votre style habituel.")
    old_course_file = st.file_uploader(
        "Importer un ancien support", 
        type=["pptx", "pdf", "txt"],
        help="L'IA analysera le vocabulaire et le format de cet ancien cours."
    )
    old_course_text_input = st.text_area(
        "Ou collez un extrait de texte",
        placeholder="Ex : Notes sur un cours précédent...",
        height=90
    )

# --- FORMULAIRE PRINCIPAL (DANS UNE CARTE CENTRÉE) ---
with st.container():
    st.markdown('<div class="custom-card">', unsafe_allow_html=True)
    st.markdown("### 📝 Paramètres de la leçon")
    
    custom_title = st.text_input("Titre du cours (optionnel)", placeholder="Ex : La découverte des 4 saisons")
    
    topic = st.text_area(
        "Sujet et notions clés *",
        placeholder="Ex : Expliquer le passage de l'automne à l'hiver, la chute des feuilles, les vêtements chauds...",
        height=110
    )
    
    c1, c2 = st.columns([1, 1])
    with c1:
        nb_slides = st.slider("Nombre de diapositives", min_value=3, max_value=8, value=4)
    with c2:
        pedagogy_notes = st.text_input(
            "Besoins spécifiques / Consignes",
            placeholder="Ex : Phrases courtes, émojis concrets, non-lecteurs..."
        )
    
    st.markdown('</div>', unsafe_allow_html=True)

# Extraction de texte du fichier modèle
def extract_text_from_file(uploaded_file):
    if uploaded_file is None:
        return ""
    try:
        if uploaded_file.name.endswith(".txt"):
            return uploaded_file.getvalue().decode("utf-8")
        elif uploaded_file.name.endswith(".pdf"):
            reader = PdfReader(uploaded_file)
            return "\n".join([page.extract_text() or "" for page in reader.pages])
        elif uploaded_file.name.endswith(".pptx"):
            prs = Presentation(uploaded_file)
            text_runs = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text_runs.append(paragraph.text)
            return "\n".join(text_runs)
    except Exception:
        return ""
    return ""

response_schema = {
    "type": "OBJECT",
    "properties": {
        "title": {"type": "STRING"},
        "slides": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "slide_title": {"type": "STRING"},
                    "adapted_content": {
                        "type": "ARRAY",
                        "items": {"type": "STRING"},
                        "description": "Phrases très courtes, simples, repères concrets avec émojis pour besoins spécifiques."
                    },
                    "advanced_content": {
                        "type": "ARRAY",
                        "items": {"type": "STRING"},
                        "description": "Notions plus poussées, phrases complètes et vocabulaire enrichi."
                    }
                },
                "required": ["slide_title", "adapted_content", "advanced_content"]
            }
        }
    },
    "required": ["title", "slides"]
}

def create_presentation(data, level_key, background_file=None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    bg_bytes = io.BytesIO(background_file.getvalue()) if background_file else None

    # Slide 1 : Titre
    slide = prs.slides.add_slide(blank_layout)
    if bg_bytes:
        bg_bytes.seek(0)
        slide.shapes.add_picture(bg_bytes, 0, 0, width=prs.slide_width, height=prs.slide_height)

    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.33), Inches(3.2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = data["title"]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 41, 59)

    p2 = tf.add_paragraph()
    level_label = "Niveau Adapté / Découverte" if level_key == "adapted_content" else "Niveau Approfondi"
    p2.text = f"— {level_label} —"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(100, 116, 139)

    # Slides suivantes
    for item in data["slides"]:
        s = prs.slides.add_slide(blank_layout)
        if bg_bytes:
            bg_bytes.seek(0)
            s.shapes.add_picture(bg_bytes, 0, 0, width=prs.slide_width, height=prs.slide_height)

        box = s.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11), Inches(5.2))
        tf_s = box.text_frame
        tf_s.word_wrap = True

        p_title = tf_s.paragraphs[0]
        p_title.text = item["slide_title"]
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(15, 23, 42)
        p_title.space_after = Pt(20)

        for bullet in item[level_key]:
            p_bullet = tf_s.add_paragraph()
            p_bullet.text = bullet
            p_bullet.font.size = Pt(22)
            p_bullet.font.color.rgb = RGBColor(51, 65, 85)
            p_bullet.space_after = Pt(14)

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out

# --- BOUTON D'ACTION ---
generate_btn = st.button("✨ Créer mes deux présentations", type="primary")

if generate_btn:
    if not api_key:
        st.error("⚠️ La clé GEMINI_API_KEY est introuvable dans les Secrets de Streamlit.")
    elif not topic.strip():
        st.warning("Veuillez indiquer un sujet pour le cours.")
    else:
        try:
            with st.spinner("Conception pédagogique différenciée en cours..."):
                extracted_reference = extract_text_from_file(old_course_file)
                combined_reference = f"{extracted_reference}\n{old_course_text_input}".strip()

                reference_context = ""
                if combined_reference:
                    reference_context = f"\n\nExemple de cours précédent à utiliser comme modèle :\n\"\"\"\n{combined_reference[:4000]}\n\"\"\""

                title_instruction = f"Titre imposé : {custom_title}" if custom_title.strip() else "Génère un titre clair et adapté."
                specific_needs = f"Besoins spécifiques : {pedagogy_notes}" if pedagogy_notes.strip() else ""

                system_prompt = (
                    "Tu es un expert en pédagogie spécialisée et en différenciation d'enseignement. "
                    "Tu dois produire deux versions distinctes pour chaque diapositive :\n"
                    "1. 'adapted_content' : syntaxe directe, phrases très courtes, repères et émojis concrets, charge cognitive allégée.\n"
                    "2. 'advanced_content' : notions complètes, phrases élaborées, vocabulaire approfondi.\n"
                    f"{reference_context}"
                )

                prompt = (
                    f"{title_instruction}\n"
                    f"Sujet : {topic}\n"
                    f"{specific_needs}\n"
                    f"Génère exactement {nb_slides} diapositives de contenu."
                )

                client = genai.Client(api_key=api_key)
                response = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        system_instruction=system_prompt,
                        response_mime_type="application/json",
                        response_schema=response_schema,
                        temperature=0.3
                    )
                )

                st.session_state["course_data"] = json.loads(response.text)
                st.success("🎉 Présentations générées avec succès !")
        except Exception as e:
            st.error(f"Une erreur est survenue : {e}")

# --- AFFICHAGE DES RÉSULTATS ---
if "course_data" in st.session_state:
    data = st.session_state["course_data"]
    
    st.markdown(f"<h2 style='text-align: center; margin-top: 1.5rem;'>📖 {data['title']}</h2>", unsafe_allow_html=True)
    
    tab1, tab2 = st.tabs(["🟢 Version Adaptée (Spécialisée)", "🔵 Version Approfondie"])

    with tab1:
        for idx, s in enumerate(data["slides"], 1):
            st.markdown(f"""
            <div class="slide-preview-card" style="border-left-color: #10b981;">
                <div class="slide-title-tag">Diapositive {idx} : {s['slide_title']}</div>
                {''.join([f'<div style=\"margin-bottom: 4px;\">• {pt}</div>' for pt in s['adapted_content']])}
            </div>
            """, unsafe_allow_html=True)
            
        pptx_1 = create_presentation(data, "adapted_content", bg_file)
        st.download_button(
            label="📥 Télécharger la présentation Adaptée (.pptx)",
            data=pptx_1,
            file_name="cours_adapte.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            key="dl_1"
        )

    with tab2:
        for idx, s in enumerate(data["slides"], 1):
            st.markdown(f"""
            <div class="slide-preview-card" style="border-left-color: #3b82f6;">
                <div class="slide-title-tag">Diapositive {idx} : {s['slide_title']}</div>
                {''.join([f'<div style=\"margin-bottom: 4px;\">• {pt}</div>' for pt in s['advanced_content']])}
            </div>
            """, unsafe_allow_html=True)
            
        pptx_2 = create_presentation(data, "advanced_content", bg_file)
        st.download_button(
            label="📥 Télécharger la présentation Approfondie (.pptx)",
            data=pptx_2,
            file_name="cours_avance.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            key="dl_2"
        )
